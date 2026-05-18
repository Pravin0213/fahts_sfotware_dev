"""Tests for Phase 4.7 temperature-results export functions in results_writer.py."""
from __future__ import annotations

import csv
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pytest

from fahts.core.results.temperature_field import TemperatureField
from fahts.core.io.results_writer import (
    export_peak_temperature_csv,
    export_temperature_history_csv,
    export_results_vtk,
    export_results_excel,
)


# ── Fixtures ──────────────────────────────────────────────────────────────────

def _make_result(n_elems: int = 3, n_steps: int = 5) -> TemperatureField:
    """Build a minimal TemperatureField for testing."""
    rng = np.random.default_rng(0)
    times = np.linspace(0.0, 3600.0, n_steps)
    eids = list(range(1, n_elems + 1))
    T_cen = rng.uniform(20.0, 700.0, (n_steps, n_elems))
    return TemperatureField(times=times, element_ids=eids, T_centroid=T_cen)


def _make_model_ns(eids: list[int]):
    """Minimal namespace model with Node + BeamElement stubs."""
    nodes = {
        1: SimpleNamespace(xyz=np.array([0.0, 0.0, 0.0])),
        2: SimpleNamespace(xyz=np.array([1.0, 0.0, 0.0])),
    }
    elements = {
        eid: SimpleNamespace(
            eid=eid, n1=1, n2=2, mat_id=1, geom_id=1,
            ecc1=None, ecc2=None,
        )
        for eid in eids
    }

    def centroid():
        return np.array([0.5, 0.0, 0.0])

    return SimpleNamespace(nodes=nodes, elements=elements, centroid=centroid)


# ── export_peak_temperature_csv ───────────────────────────────────────────────

class TestExportPeakTemperatureCsv:
    def test_file_created(self, tmp_path):
        result = _make_result()
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, out)
        assert out.exists()

    def test_has_header(self, tmp_path):
        result = _make_result()
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, out)
        with out.open() as fh:
            reader = csv.DictReader(fh)
            assert "element_id" in reader.fieldnames
            assert "T_peak_centroid_C" in reader.fieldnames

    def test_row_count_matches_elements(self, tmp_path):
        n = 4
        result = _make_result(n_elems=n)
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, out)
        with out.open() as fh:
            rows = list(csv.DictReader(fh))
        assert len(rows) == n

    def test_element_ids_in_csv(self, tmp_path):
        result = _make_result(n_elems=3)
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, out)
        with out.open() as fh:
            rows = list(csv.DictReader(fh))
        ids = [int(r["element_id"]) for r in rows]
        assert ids == [1, 2, 3]

    def test_passes_600_column_present(self, tmp_path):
        result = _make_result()
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, out)
        with out.open() as fh:
            reader = csv.DictReader(fh)
            assert "passes_600" in reader.fieldnames

    def test_accepts_path_string(self, tmp_path):
        result = _make_result()
        out = tmp_path / "peak.csv"
        export_peak_temperature_csv(result, str(out))
        assert out.exists()


# ── export_temperature_history_csv ────────────────────────────────────────────

class TestExportTemperatureHistoryCsv:
    def test_file_created(self, tmp_path):
        result = _make_result()
        out = tmp_path / "history.csv"
        export_temperature_history_csv(result, out)
        assert out.exists()

    def test_has_time_columns(self, tmp_path):
        result = _make_result()
        out = tmp_path / "history.csv"
        export_temperature_history_csv(result, out)
        with out.open() as fh:
            reader = csv.DictReader(fh)
            assert "time_s" in reader.fieldnames
            assert "time_min" in reader.fieldnames

    def test_row_count_matches_steps(self, tmp_path):
        n_steps = 7
        result = _make_result(n_steps=n_steps)
        out = tmp_path / "history.csv"
        export_temperature_history_csv(result, out)
        with out.open() as fh:
            rows = list(csv.DictReader(fh))
        assert len(rows) == n_steps

    def test_element_columns_present(self, tmp_path):
        result = _make_result(n_elems=3)
        out = tmp_path / "history.csv"
        export_temperature_history_csv(result, out)
        with out.open() as fh:
            reader = csv.DictReader(fh)
            names = reader.fieldnames
        assert "T_eid_1" in names
        assert "T_eid_2" in names
        assert "T_eid_3" in names

    def test_accepts_path_string(self, tmp_path):
        result = _make_result()
        out = tmp_path / "history.csv"
        export_temperature_history_csv(result, str(out))
        assert out.exists()


# ── export_results_vtk ────────────────────────────────────────────────────────

class TestExportResultsVtk:
    def test_file_created(self, tmp_path):
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        assert out.exists()

    def test_vtk_can_be_read_back(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert mesh is not None

    def test_vtk_cell_count(self, tmp_path):
        import pyvista as pv
        n = 3
        result = _make_result(n_elems=n)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert mesh.n_cells == n

    def test_vtk_has_element_id_array(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "element_id" in mesh.cell_data

    def test_vtk_has_T_peak_array(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "T_peak_C" in mesh.cell_data

    def test_vtk_has_T_final_array(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "T_final_C" in mesh.cell_data

    def test_vtk_has_fire_resistance_array(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "fire_resistance_min" in mesh.cell_data

    def test_vtk_has_passes_600_array(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "passes_600" in mesh.cell_data

    def test_vtk_has_per_step_arrays(self, tmp_path):
        import pyvista as pv
        n_steps = 4
        result = _make_result(n_elems=2, n_steps=n_steps)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        assert "T_step000_C" in mesh.cell_data
        assert f"T_step{n_steps - 1:03d}_C" in mesh.cell_data

    def test_vtk_element_ids_correct(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=3)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        ids = sorted(mesh.cell_data["element_id"].tolist())
        assert ids == sorted(result.element_ids)

    def test_vtk_T_peak_values_correct(self, tmp_path):
        import pyvista as pv
        result = _make_result(n_elems=2)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        mesh = pv.read(str(out))
        # Reconstruct expected: sort by element_id
        eid_arr = mesh.cell_data["element_id"]
        T_peak_arr = mesh.cell_data["T_peak_C"]
        for i, eid in enumerate(eid_arr):
            expected = result.peak_centroid_temperature(int(eid))
            assert abs(T_peak_arr[i] - expected) < 1e-4

    def test_vtk_no_crash_when_element_missing_from_model(self, tmp_path):
        """Extra eids in result not present in model are silently skipped."""
        result = _make_result(n_elems=3)   # eids 1, 2, 3
        model = _make_model_ns([1, 2])      # only 1 and 2 in model
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, out)
        assert out.exists()

    def test_vtk_accepts_path_string(self, tmp_path):
        result = _make_result(n_elems=1)
        model = _make_model_ns(result.element_ids)
        out = tmp_path / "results.vtk"
        export_results_vtk(result, model, str(out))
        assert out.exists()


# ── export_results_excel ─────────────────────────────────────────────────────

class TestExportResultsExcel:
    def test_file_created(self, tmp_path):
        pytest.importorskip("openpyxl")
        result = _make_result()
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        assert out.exists()

    def test_has_summary_sheet(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        result = _make_result()
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        wb = openpyxl.load_workbook(out)
        assert "Summary" in wb.sheetnames

    def test_has_t_history_sheet(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        result = _make_result()
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        wb = openpyxl.load_workbook(out)
        assert "T_History" in wb.sheetnames

    def test_summary_row_count(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        n = 5
        result = _make_result(n_elems=n)
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Summary"]
        # 1 header row + n data rows
        assert ws.max_row == n + 1

    def test_history_row_count(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        n_steps = 6
        result = _make_result(n_steps=n_steps)
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["T_History"]
        assert ws.max_row == n_steps + 1

    def test_summary_has_element_id_column(self, tmp_path):
        pytest.importorskip("openpyxl")
        import openpyxl
        result = _make_result()
        out = tmp_path / "results.xlsx"
        export_results_excel(result, out)
        wb = openpyxl.load_workbook(out)
        ws = wb["Summary"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        assert "element_id" in headers

    def test_accepts_path_string(self, tmp_path):
        pytest.importorskip("openpyxl")
        result = _make_result()
        out = tmp_path / "results.xlsx"
        export_results_excel(result, str(out))
        assert out.exists()
